from __future__ import annotations

from typing import Any

from fastapi import HTTPException
from sqlalchemy import text
from sqlalchemy.engine import Connection

from app.db.session import get_connection, get_transaction
from app.services.inventory_control_service import (
    list_producto_inventario_config,
    normalize_clasificacion_inventario,
    replace_producto_inventario_config,
)
from app.utils.company_access import (
    build_in_clause,
    can_access_company,
    normalize_company_ids,
    user_company_ids,
)


def build_inventory_config_payload(
    inventory_config: list[dict[str, Any]] | None,
    companies: list[int],
    body: dict[str, Any],
) -> list[dict[str, Any]] | None:
    if isinstance(inventory_config, list):
        return inventory_config

    has_inline_config = body.get("Almacen_Id") or body.get("ClasificacionInventario")
    if not has_inline_config:
        return None

    company_ids = companies or [
        company_id
        for company_id in [
            body.get("Company_Id"),
            body.get("company_id"),
        ]
        if company_id
    ]
    if not company_ids:
        return None

    return [
        {
            "Company_Id": int(company_id),
            "Almacen_Id": int(body["Almacen_Id"]) if body.get("Almacen_Id") else None,
            "ClasificacionInventario": normalize_clasificacion_inventario(
                body.get("ClasificacionInventario")
            ),
            "Activo": False if body.get("ConfiguracionInventarioActiva") is False else True,
        }
        for company_id in company_ids
    ]


def _validate_product_company_access(
    connection: Connection,
    product_id: int,
    current_user: dict[str, Any],
) -> None:
    if current_user.get("is_admin"):
        return
    allowed = set(user_company_ids(current_user))
    if not allowed:
        raise HTTPException(status_code=404, detail="Producto no encontrado")

    rows = connection.execute(
        text(
            """
            SELECT Company_Id
            FROM ERP_PRODUCTO_EMPRESA
            WHERE Producto_Id = :product_id
            """
        ),
        {"product_id": product_id},
    ).mappings().all()
    product_companies = {int(row["Company_Id"]) for row in rows}
    if product_companies and not product_companies.intersection(allowed):
        raise HTTPException(status_code=404, detail="Producto no encontrado")


def _validate_sat_keys(
    connection: Connection,
    clave_prodserv: str,
    clave_unidad: str,
) -> None:
    prodserv = connection.execute(
        text("SELECT Clave FROM SAT_CLAVE_PRODSERV WHERE Clave = :clave"),
        {"clave": clave_prodserv},
    ).mappings().first()
    if not prodserv:
        raise HTTPException(
            status_code=400,
            detail=f"ClaveProdServSAT {clave_prodserv} no existe en catalogo SAT",
        )

    unidad = connection.execute(
        text("SELECT Clave FROM SAT_UNIDADES WHERE Clave = :clave"),
        {"clave": clave_unidad},
    ).mappings().first()
    if not unidad:
        raise HTTPException(
            status_code=400,
            detail=f"ClaveUnidadSAT {clave_unidad} no existe en catalogo SAT",
        )


def _normalize_companies(
    current_user: dict[str, Any],
    companies: list[int] | None,
) -> list[int]:
    normalized = normalize_company_ids(companies)
    if current_user.get("is_admin"):
        return normalized

    allowed = set(user_company_ids(current_user))
    if not allowed:
        raise HTTPException(status_code=403, detail="No tiene empresas asignadas")

    if not normalized:
        return list(allowed)

    if not set(normalized).issubset(allowed):
        raise HTTPException(
            status_code=403,
            detail="Solo puede usar empresas a las que pertenece",
        )
    return normalized


def list_products(
    current_user: dict[str, Any],
    page: int | None,
    limit: int | None,
    search: str,
    activo: str | None,
    company_id: str | None,
) -> dict[str, Any]:
    conditions = ["1 = 1"]
    params: dict[str, Any] = {}

    if search:
        conditions.append("(p.SKU LIKE :search OR p.Nombre LIKE :search OR p.Descripcion LIKE :search)")
        params["search"] = f"%{search}%"

    if activo is not None:
        conditions.append("p.Activo = :activo")
        params["activo"] = 1 if str(activo).lower() == "true" else 0
    else:
        conditions.append("p.Activo = 1")

    if not current_user.get("is_admin"):
        companies = user_company_ids(current_user)
        if not companies:
            if page and limit:
                page = max(int(page), 1)
                limit = max(int(limit), 1)
                return {
                    "data": [],
                    "pagination": {
                        "total": 0,
                        "page": page,
                        "limit": limit,
                        "pages": 0,
                    },
                }
            return {"data": [], "total": 0}
        clause, clause_params = build_in_clause("company", companies)
        conditions.append(
            f"""
            EXISTS (
                SELECT 1
                FROM ERP_PRODUCTO_EMPRESA pe
                WHERE pe.Producto_Id = p.Producto_Id
                  AND pe.Company_Id IN ({clause})
            )
            """
        )
        params.update(clause_params)
    elif company_id and company_id != "all":
        conditions.append(
            """
            EXISTS (
                SELECT 1
                FROM ERP_PRODUCTO_EMPRESA pe
                WHERE pe.Producto_Id = p.Producto_Id
                  AND pe.Company_Id = :company_id
            )
            """
        )
        params["company_id"] = int(company_id)

    where_clause = " AND ".join(conditions)
    count_query = f"SELECT COUNT(*) AS total FROM ERP_PRODUCTOS p WHERE {where_clause}"
    data_query = f"SELECT p.* FROM ERP_PRODUCTOS p WHERE {where_clause} ORDER BY p.Producto_Id DESC"

    if page and limit:
        page = max(int(page), 1)
        limit = max(int(limit), 1)
        params["offset"] = (page - 1) * limit
        params["limit"] = limit
        data_query += " OFFSET :offset ROWS FETCH NEXT :limit ROWS ONLY"

    with get_connection() as connection:
        total_row = connection.execute(text(count_query), params).mappings().first()
        result = connection.execute(text(data_query), params)
        rows = [dict(row) for row in result.mappings().all()]

    response: dict[str, Any] = {"data": rows}
    total = int(total_row["total"] if total_row else 0)
    if page and limit:
        response["pagination"] = {
            "total": total,
            "page": page,
            "limit": limit,
            "pages": (total + limit - 1) // limit,
        }
    else:
        response["total"] = total
    return response


def get_product(product_id: int, current_user: dict[str, Any]) -> dict[str, Any]:
    with get_connection() as connection:
        _validate_product_company_access(connection, product_id, current_user)
        product = connection.execute(
            text("SELECT * FROM ERP_PRODUCTOS WHERE Producto_Id = :product_id"),
            {"product_id": product_id},
        ).mappings().first()
        if not product:
            raise HTTPException(status_code=404, detail="Producto no encontrado")

        companies = connection.execute(
            text(
                """
                SELECT c.Company_Id, c.NameCompany
                FROM ERP_PRODUCTO_EMPRESA pe
                INNER JOIN ERP_COMPANY c ON pe.Company_Id = c.Company_Id
                WHERE pe.Producto_Id = :product_id
                """
            ),
            {"product_id": product_id},
        ).mappings().all()
        inventory_config = list_producto_inventario_config(product_id, connection)

    response = dict(product)
    response["companies"] = [dict(row) for row in companies]
    response["inventoryConfig"] = inventory_config
    return response


def create_product(payload: dict[str, Any], current_user: dict[str, Any]) -> dict[str, Any]:
    required_fields = ["SKU", "Nombre", "ClaveProdServSAT", "ClaveUnidadSAT"]
    missing = [field for field in required_fields if not payload.get(field)]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Campos requeridos: {', '.join(required_fields)}",
        )

    companies = _normalize_companies(current_user, payload.get("companies"))
    config_payload = build_inventory_config_payload(
        payload.get("inventoryConfig"),
        companies,
        payload,
    )

    try:
        with get_transaction() as connection:
            _validate_sat_keys(
                connection,
                str(payload["ClaveProdServSAT"]),
                str(payload["ClaveUnidadSAT"]),
            )

            result = connection.execute(
                text(
                    """
                    INSERT INTO ERP_PRODUCTOS (
                        SKU,
                        Nombre,
                        Descripcion,
                        Precio,
                        TipoMoneda,
                        ClaveProdServSAT,
                        ClaveUnidadSAT,
                        ImpuestoIVA,
                        Activo,
                        CreadoPor
                    )
                    OUTPUT INSERTED.Producto_Id
                    VALUES (
                        :sku,
                        :nombre,
                        :descripcion,
                        :precio,
                        :tipo_moneda,
                        :clave_prodserv_sat,
                        :clave_unidad_sat,
                        :impuesto_iva,
                        :activo,
                        :creado_por
                    )
                    """
                ),
                {
                    "sku": payload["SKU"],
                    "nombre": payload["Nombre"],
                    "descripcion": payload.get("Descripcion"),
                    "precio": payload.get("Precio") or 0,
                    "tipo_moneda": payload.get("TipoMoneda"),
                    "clave_prodserv_sat": payload["ClaveProdServSAT"],
                    "clave_unidad_sat": payload["ClaveUnidadSAT"],
                    "impuesto_iva": payload.get("ImpuestoIVA") or 16.0,
                    "activo": False if payload.get("Activo") is False else True,
                    "creado_por": current_user.get("User_Id"),
                },
            )
            row = result.first()
            product_id = int(row[0]) if row else None

            for company_id in companies:
                connection.execute(
                    text(
                        """
                        INSERT INTO ERP_PRODUCTO_EMPRESA (Producto_Id, Company_Id)
                        VALUES (:product_id, :company_id)
                        """
                    ),
                    {"product_id": product_id, "company_id": company_id},
                )

            if config_payload:
                replace_producto_inventario_config(
                    connection,
                    product_id,
                    config_payload,
                    str(current_user.get("Username") or current_user.get("Email") or ""),
                )
    except HTTPException:
        raise
    except Exception as exc:
        message = str(exc).lower()
        if "duplicate" in message or "2627" in message:
            raise HTTPException(status_code=409, detail="El SKU ya existe") from exc
        raise

    return {"msg": "Producto creado", "Producto_Id": product_id}


def update_product(
    product_id: int,
    payload: dict[str, Any],
    current_user: dict[str, Any],
) -> dict[str, str]:
    companies_payload = payload.get("companies")
    companies = (
        _normalize_companies(current_user, companies_payload)
        if companies_payload is not None
        else None
    )
    config_payload = build_inventory_config_payload(
        payload.get("inventoryConfig"),
        companies or [],
        payload,
    )

    with get_transaction() as connection:
        _validate_product_company_access(connection, product_id, current_user)

        current_row = connection.execute(
            text(
                """
                SELECT ClaveProdServSAT, ClaveUnidadSAT
                FROM ERP_PRODUCTOS
                WHERE Producto_Id = :product_id
                """
            ),
            {"product_id": product_id},
        ).mappings().first()
        if not current_row:
            raise HTTPException(status_code=404, detail="Producto no encontrado")

        if payload.get("ClaveProdServSAT") or payload.get("ClaveUnidadSAT"):
            _validate_sat_keys(
                connection,
                str(payload.get("ClaveProdServSAT") or current_row["ClaveProdServSAT"]),
                str(payload.get("ClaveUnidadSAT") or current_row["ClaveUnidadSAT"]),
            )

        connection.execute(
            text(
                """
                UPDATE ERP_PRODUCTOS
                SET
                    Nombre = COALESCE(:nombre, Nombre),
                    Descripcion = :descripcion,
                    Precio = COALESCE(:precio, Precio),
                    TipoMoneda = :tipo_moneda,
                    ClaveProdServSAT = COALESCE(:clave_prodserv_sat, ClaveProdServSAT),
                    ClaveUnidadSAT = COALESCE(:clave_unidad_sat, ClaveUnidadSAT),
                    ImpuestoIVA = COALESCE(:impuesto_iva, ImpuestoIVA),
                    Activo = COALESCE(:activo, Activo),
                    FechaActualizacion = GETDATE(),
                    ActualizadoPor = :actualizado_por
                WHERE Producto_Id = :product_id
                """
            ),
            {
                "product_id": product_id,
                "nombre": payload.get("Nombre"),
                "descripcion": payload.get("Descripcion"),
                "precio": payload.get("Precio"),
                "tipo_moneda": payload.get("TipoMoneda"),
                "clave_prodserv_sat": payload.get("ClaveProdServSAT"),
                "clave_unidad_sat": payload.get("ClaveUnidadSAT"),
                "impuesto_iva": payload.get("ImpuestoIVA"),
                "activo": payload.get("Activo"),
                "actualizado_por": current_user.get("User_Id"),
            },
        )

        if companies is not None:
            connection.execute(
                text("DELETE FROM ERP_PRODUCTO_EMPRESA WHERE Producto_Id = :product_id"),
                {"product_id": product_id},
            )
            for company_id in companies:
                connection.execute(
                    text(
                        """
                        INSERT INTO ERP_PRODUCTO_EMPRESA (Producto_Id, Company_Id)
                        VALUES (:product_id, :company_id)
                        """
                    ),
                    {"product_id": product_id, "company_id": company_id},
                )

        if config_payload is not None:
            replace_producto_inventario_config(
                connection,
                product_id,
                config_payload,
                str(current_user.get("Username") or current_user.get("Email") or ""),
            )
        elif companies is not None:
            existing_config = list_producto_inventario_config(product_id, connection)
            allowed_companies = set(companies)
            filtered = [
                config
                for config in existing_config
                if int(config["Company_Id"]) in allowed_companies
            ]
            replace_producto_inventario_config(
                connection,
                product_id,
                filtered,
                str(current_user.get("Username") or current_user.get("Email") or ""),
            )

    return {"msg": "Producto actualizado"}


def delete_product(product_id: int, current_user: dict[str, Any]) -> dict[str, str]:
    child_deletes = [
        ("ERP_SOLICITUD_PRECIO_DETALLE", "Producto_Id"),
        ("ERP_SOLICITUD_CAMBIO_PRECIO", "Producto_Id"),
        ("ERP_PRECIOS_CLIENTE_PRODUCTO", "Producto_Id"),
        ("ERP_CLIENT_RECURRING_PRODUCTS", "Producto_Id"),
        ("ERP_COTIZACION_DETALLE", "ID_PRODUCTO"),
        ("ERP_KARDEX", "Producto_Id"),
        ("ERP_STOCK", "Producto_Id"),
        ("ERP_PRODUCTO_ALMACEN_CONFIG", "Producto_Id"),
        ("ERP_PRODUCTO_EMPRESA", "Producto_Id"),
        ("ERP_CALIDAD_ALERTA", "Producto_Id"),
        ("ERP_COMPRA_ORDEN_DETALLE", "Producto_Id"),
        ("ERP_COMPRA_RECEPCION_DETALLE", "Producto_Id"),
        ("ERP_COSTEO_MENSUAL", "Producto_Id"),
        ("ERP_INVENTARIO_ESTADO_PRODUCTO", "Producto_Id"),
        ("ERP_RECEPCION_PRODUCTO_TERMINADO", "Producto_Id"),
    ]

    with get_transaction() as connection:
        _validate_product_company_access(connection, product_id, current_user)
        for table_name, column_name in child_deletes:
            connection.execute(
                text(f"DELETE FROM {table_name} WHERE {column_name} = :product_id"),
                {"product_id": product_id},
            )
        connection.execute(
            text("DELETE FROM ERP_PRODUCTOS WHERE Producto_Id = :product_id"),
            {"product_id": product_id},
        )
    return {"msg": "Producto eliminado"}


def import_products_not_implemented() -> dict[str, str]:
    return {
        "msg": "Importacion pendiente",
        "detail": (
            "La importacion avanzada desde Excel sigue pendiente de portarse a Python. "
            "El CRUD de productos e inventario ya esta activo."
        ),
    }


# ── Header map (normalised key → canonical column name) ──────────────────────

_HEADER_MAP: dict[str, str] = {
    "sku": "SKU", "codigo": "SKU", "codigoproducto": "SKU", "clave": "SKU",
    "nombre": "Nombre", "nombreproducto": "Nombre",
    "descripcion": "Descripcion",
    "precio": "Precio", "preciounitario": "Precio", "precioventa": "Precio", "preciove": "Precio",
    "tipomoneda": "TipoMoneda", "moneda": "TipoMoneda", "tipodemoneda": "TipoMoneda",
    "currency": "TipoMoneda", "divisa": "TipoMoneda",
    "claveprodservsat": "ClaveProdServSAT", "claveprodserv": "ClaveProdServSAT",
    "claveproducto": "ClaveProdServSAT", "clavesat": "ClaveProdServSAT",
    "claveunidadsat": "ClaveUnidadSAT", "claveunidad": "ClaveUnidadSAT",
    "unidadsat": "ClaveUnidadSAT", "unidadde": "ClaveUnidadSAT", "unidad": "ClaveUnidadSAT",
    "impuestoiva": "ImpuestoIVA", "iva": "ImpuestoIVA",
    "activo": "Activo", "estatus": "Activo", "status": "Activo",
    "almacen_id": "Almacen_Id", "almacenid": "Almacen_Id",
    "clasificacioninventario": "ClasificacionInventario",
}

_MONEDA_MAP: dict[str, str] = {
    "MXN": "MXN", "PESOS": "MXN", "PESO": "MXN", "PESO MEXICANO": "MXN",
    "PESOS MEXICANOS": "MXN", "MX": "MXN", "M.N.": "MXN", "MN": "MXN",
    "USD": "USD", "DOLAR": "USD", "DOLARES": "USD", "DÓLAR": "USD",
    "DÓLARES": "USD", "DOLLAR": "USD", "DOLLARS": "USD", "US": "USD",
    "EUR": "EUR", "EURO": "EUR", "EUROS": "EUR",
}


def _normalize_key(key: str) -> str:
    import unicodedata, re
    nfkd = unicodedata.normalize("NFD", str(key))
    without_accents = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-zA-Z0-9]", "", without_accents).lower()


def import_products(file_bytes: bytes, filename: str, company_id: int, user_id: int | None) -> dict[str, Any]:
    import openpyxl
    import io, re
    from app.services.inventory_control_service import upsert_producto_inventario_config

    # ── Parse Excel ──────────────────────────────────────────────────────────
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    header_raw = [str(c) if c is not None else "" for c in rows[0]]
    header_norm = [_normalize_key(h) for h in header_raw]

    data: list[dict[str, Any]] = []
    for row in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        record: dict[str, Any] = {}
        for i, val in enumerate(row):
            if i >= len(header_norm):
                break
            nk = header_norm[i]
            target = _HEADER_MAP.get(nk)
            if not target:
                if "unidad" in nk:
                    target = "ClaveUnidadSAT"
                elif "claveprodserv" in nk or "clavesat" in nk:
                    target = "ClaveProdServSAT"
            if target:
                record[target] = val
        data.append(record)

    if not data:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    # ── Load SAT catalogues ──────────────────────────────────────────────────
    with get_connection() as conn:
        prodserv_rows = conn.execute(text("SELECT Clave FROM SAT_CLAVE_PRODSERV")).fetchall()
        unidades_rows = conn.execute(text("SELECT Clave, Nombre, Simbolo FROM SAT_UNIDADES")).fetchall()

    sat_prodserv: set[str] = {str(r[0]).strip().upper() for r in prodserv_rows}
    sat_unidades: set[str] = {str(r[0]).strip().upper() for r in unidades_rows}
    sat_by_simbolo: dict[str, str] = {}
    sat_by_nombre: dict[str, str] = {}
    for r in unidades_rows:
        clave = str(r[0]).strip().upper()
        if r[2]:
            sat_by_simbolo[str(r[2]).strip().upper()] = clave
        if r[1]:
            sat_by_nombre[str(r[1]).strip().upper()] = clave

    # ── Pre-load existing SKUs ───────────────────────────────────────────────
    sku_values = {str(row.get("SKU", "")) for row in data if row.get("SKU")}
    existing_skus: set[str] = set()
    if sku_values:
        placeholders = ", ".join(f":sku{i}" for i in range(len(sku_values)))
        params = {f"sku{i}": v for i, v in enumerate(sku_values)}
        with get_connection() as conn:
            result = conn.execute(
                text(f"SELECT SKU FROM ERP_PRODUCTOS WHERE SKU IN ({placeholders})"),
                params,
            ).fetchall()
        existing_skus = {str(r[0]) for r in result}

    errores: list[dict] = []
    exitosas = 0
    con_error = 0

    for idx, row in enumerate(data):
        fila = idx + 2

        # ── Derive Nombre ────────────────────────────────────────────────────
        if not row.get("Nombre"):
            row["Nombre"] = str(row.get("Descripcion", "") or row.get("SKU", "") or "")[:200] or None

        sku = row.get("SKU")
        nombre = row.get("Nombre")
        descripcion = row.get("Descripcion")
        precio = row.get("Precio")
        tipo_moneda = row.get("TipoMoneda")
        clave_prod = row.get("ClaveProdServSAT")
        clave_unidad = row.get("ClaveUnidadSAT")
        iva = row.get("ImpuestoIVA")
        activo = row.get("Activo")

        # ── Normalise moneda ─────────────────────────────────────────────────
        tipo_moneda = _MONEDA_MAP.get(str(tipo_moneda or "").strip().upper(), "MXN")

        # ── Required field validation ────────────────────────────────────────
        missing = [f for f, v in [("SKU", sku), ("Nombre", nombre), ("ClaveProdServSAT", clave_prod), ("ClaveUnidadSAT", clave_unidad)] if not v]
        if missing:
            errores.append({"fila": fila, "error": f"Campos requeridos faltantes: {', '.join(missing)}", "datos": row})
            con_error += 1
            continue

        # ── Validate ClaveProdServSAT ────────────────────────────────────────
        clave_prod_norm = str(clave_prod).strip().upper()
        if clave_prod_norm not in sat_prodserv:
            texto = f"{nombre or ''} {descripcion or ''}".strip()
            if texto:
                try:
                    with get_connection() as conn:
                        alt = conn.execute(
                            text("SELECT TOP 1 Clave FROM SAT_CLAVE_PRODSERV WHERE Descripcion LIKE :s ORDER BY Clave"),
                            {"s": f"%{texto[:40]}%"},
                        ).fetchone()
                    if alt:
                        clave_prod_norm = str(alt[0]).strip().upper()
                        clave_prod = clave_prod_norm
                        sat_prodserv.add(clave_prod_norm)
                except Exception:
                    pass

        if clave_prod_norm not in sat_prodserv:
            errores.append({"fila": fila, "error": f"ClaveProdServSAT {clave_prod} no existe en catálogo SAT", "datos": row})
            con_error += 1
            continue

        # ── Validate ClaveUnidadSAT ──────────────────────────────────────────
        clave_unidad_norm = str(clave_unidad).strip().upper()
        if clave_unidad_norm == "PZ":
            clave_unidad_norm = "PZA"

        if clave_unidad_norm not in sat_unidades:
            mapped = sat_by_simbolo.get(clave_unidad_norm) or sat_by_nombre.get(clave_unidad_norm)
            if not mapped:
                texto = f"{nombre or ''} {descripcion or ''}".lower()
                if re.search(r"kilo|kilogram|kg\b", texto):
                    mapped = sat_by_simbolo.get("KG") or sat_by_nombre.get("KILOGRAMO") or "KGM"
                elif re.search(r"litro|lt\b|ltr\b", texto):
                    mapped = sat_by_nombre.get("LITRO") or "LTR"
                elif re.search(r"pieza|pza", texto):
                    mapped = sat_by_simbolo.get("PZA") or sat_by_nombre.get("PIEZA") or "H87"
                elif re.search(r"servicio", texto):
                    mapped = sat_by_nombre.get("UNIDAD DE SERVICIO") or "E48"
                elif not mapped:
                    mapped = "H87" if "H87" in sat_unidades else ("E48" if "E48" in sat_unidades else None)
            if mapped and str(mapped).upper() in sat_unidades:
                clave_unidad_norm = str(mapped).upper()
            else:
                errores.append({"fila": fila, "error": f"ClaveUnidadSAT {clave_unidad} no existe en catálogo SAT", "datos": row})
                con_error += 1
                continue

        # ── Normalise numeric / bool fields ──────────────────────────────────
        try:
            precio_f = float(str(precio).replace(",", "")) if precio not in (None, "", "None") else 0.0
        except (ValueError, TypeError):
            precio_f = 0.0
        try:
            iva_f = float(str(iva)) if iva not in (None, "", "None") else 16.0
        except (ValueError, TypeError):
            iva_f = 16.0
        activo_bit = 0 if activo in (False, 0, "0", "false", "False", "inactivo", "Inactivo") else 1

        try:
            sku_str = str(sku)
            if sku_str in existing_skus:
                with get_transaction() as conn:
                    conn.execute(
                        text("""
                            UPDATE ERP_PRODUCTOS SET
                                Nombre=:nombre, Descripcion=:desc, Precio=:precio,
                                TipoMoneda=:moneda, ClaveProdServSAT=:cps, ClaveUnidadSAT=:cus,
                                ImpuestoIVA=:iva, Activo=:activo,
                                FechaActualizacion=GETDATE(), ActualizadoPor=:uid
                            WHERE SKU=:sku
                        """),
                        {"nombre": str(nombre), "desc": str(descripcion) if descripcion else None,
                         "precio": precio_f, "moneda": tipo_moneda, "cps": clave_prod_norm,
                         "cus": clave_unidad_norm, "iva": iva_f, "activo": activo_bit,
                         "uid": user_id, "sku": sku_str},
                    )
                    conn.execute(
                        text("""
                            INSERT INTO ERP_PRODUCTO_EMPRESA (Producto_Id, Company_Id)
                            SELECT p.Producto_Id, :cid
                            FROM ERP_PRODUCTOS p
                            WHERE p.SKU = :sku
                              AND NOT EXISTS (
                                SELECT 1 FROM ERP_PRODUCTO_EMPRESA pe
                                WHERE pe.Producto_Id = p.Producto_Id AND pe.Company_Id = :cid
                              )
                        """),
                        {"cid": company_id, "sku": sku_str},
                    )
                    prod_row = conn.execute(
                        text("SELECT TOP 1 Producto_Id FROM ERP_PRODUCTOS WHERE SKU=:sku"),
                        {"sku": sku_str},
                    ).fetchone()
                prod_id = prod_row[0] if prod_row else None
                if prod_id and (row.get("Almacen_Id") or row.get("ClasificacionInventario")):
                    upsert_producto_inventario_config(prod_id, {
                        "Company_Id": company_id,
                        "Almacen_Id": int(row["Almacen_Id"]) if row.get("Almacen_Id") else None,
                        "ClasificacionInventario": row.get("ClasificacionInventario"),
                    }, None)
            else:
                with get_transaction() as conn:
                    result = conn.execute(
                        text("""
                            INSERT INTO ERP_PRODUCTOS
                                (SKU, Nombre, Descripcion, Precio, TipoMoneda, ClaveProdServSAT,
                                 ClaveUnidadSAT, ImpuestoIVA, Activo, CreadoPor)
                            OUTPUT INSERTED.Producto_Id
                            VALUES (:sku, :nombre, :desc, :precio, :moneda, :cps, :cus, :iva, :activo, :uid)
                        """),
                        {"sku": sku_str, "nombre": str(nombre),
                         "desc": str(descripcion) if descripcion else None,
                         "precio": precio_f, "moneda": tipo_moneda, "cps": clave_prod_norm,
                         "cus": clave_unidad_norm, "iva": iva_f, "activo": activo_bit,
                         "uid": user_id},
                    )
                    new_row = result.fetchone()
                    new_id = new_row[0] if new_row else None
                    if new_id:
                        conn.execute(
                            text("INSERT INTO ERP_PRODUCTO_EMPRESA (Producto_Id, Company_Id) VALUES (:pid, :cid)"),
                            {"pid": new_id, "cid": company_id},
                        )
                existing_skus.add(sku_str)
                new_id_val = new_id if "new_id" in dir() else None
                if new_id_val and (row.get("Almacen_Id") or row.get("ClasificacionInventario")):
                    upsert_producto_inventario_config(new_id_val, {
                        "Company_Id": company_id,
                        "Almacen_Id": int(row["Almacen_Id"]) if row.get("Almacen_Id") else None,
                        "ClasificacionInventario": row.get("ClasificacionInventario"),
                    }, None)
            exitosas += 1
        except Exception as exc:
            errores.append({"fila": fila, "error": str(exc), "datos": row})
            con_error += 1

    # ── Save import log ──────────────────────────────────────────────────────
    import json
    try:
        with get_transaction() as conn:
            conn.execute(
                text("""
                    INSERT INTO ERP_IMPORTACIONES_LOG
                        (NombreArchivo, TotalFilas, FilasExitosas, FilasConError, Usuario_Id, Errores)
                    VALUES (:fname, :total, :ok, :err, :uid, :errs)
                """),
                {"fname": filename, "total": len(data), "ok": exitosas,
                 "err": con_error, "uid": user_id,
                 "errs": json.dumps(errores[:200], ensure_ascii=False)},
            )
    except Exception:
        pass

    return {
        "msg": "Importación completada",
        "total": len(data),
        "exitosas": exitosas,
        "conError": con_error,
        "errores": errores[:100],
    }
